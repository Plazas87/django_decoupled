"""Services module."""
from io import BytesIO
from typing import Set

import pandas as pd
from openpyxl import load_workbook

from ...application.dtos import FileCategory, FileDocument, FileWorkspace
from ...application.interfaces import IFileReader


class ExcelFileReader(IFileReader[FileWorkspace]):
    """ExcelFileReader class."""

    def read_from_bytes(self, bytes: bytes) -> Set[FileWorkspace]:
        """Read an  file form bytes."""
        workspaces = set()
        categories = []

        bytes.seek(0)  # type: ignore
        wb = load_workbook(filename=BytesIO(bytes.read()))  # type: ignore

        for ws in wb:
            workspace_name = ws.title
            df = pd.read_excel(wb, engine="openpyxl", sheet_name=workspace_name)
            df.rename(
                columns={
                    df.columns[0]: "category",
                    df.columns[1]: "text",
                },
                inplace=True,
            )
            for category, documents_df in df.groupby(by="category"):
                documents = []

                for _document in documents_df.itertuples(name="Document", index=False):
                    documents.append(FileDocument(text=_document.text))

                categories.append(FileCategory(name=category, documents=documents))

            workspaces.add(FileWorkspace(name=workspace_name, categories=categories))

        return workspaces
